"""Exercise installed content tools inside the sandbox, writing only private temporary files."""
import json
from pathlib import Path
import subprocess
import tempfile
import zipfile

from bs4 import BeautifulSoup
from docx import Document
from lxml import etree
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter


def run(*args):
    return subprocess.run(args, check=True, capture_output=True, text=True, timeout=10).stdout


with tempfile.TemporaryDirectory(prefix='poketto-toolkit-') as directory:
    root = Path(directory)
    text = '\u6d4b\u8bd5\u56fe\u7247'
    image = Image.new('RGB', (640, 480), 'white')
    font = ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 24)
    ImageDraw.Draw(image).text((10, 10), text, fill='black', font=font)
    image.resize((320, 240)).save(root / 'image.webp')
    assert Image.open(root / 'image.webp').size == (320, 240)
    assert BeautifulSoup('<p>' + text + '</p>', 'lxml').p.text == text
    assert etree.fromstring(('<value>' + text + '</value>').encode()).text == text
    document = Document()
    document.add_paragraph(text)
    document.save(root / 'text.docx')
    assert Document(root / 'text.docx').paragraphs[0].text == text
    workbook = Workbook()
    workbook.active['A1'] = text
    workbook.save(root / 'book.xlsx')
    assert load_workbook(root / 'book.xlsx').active['A1'].value == text
    image.save(root / 'page.pdf')
    reader, writer = PdfReader(root / 'page.pdf'), PdfWriter()
    assert len(reader.pages) == 1
    writer.add_page(reader.pages[0])
    writer.write(root / 'copy.pdf')
    assert len(PdfReader(root / 'copy.pdf').pages) == 1
    assert 'Pages:' in run('pdfinfo', str(root / 'page.pdf'))
    run('pdftotext', str(root / 'page.pdf'), str(root / 'page.txt'))
    run('pdftoppm', '-singlefile', '-scale-to', '200', '-png', str(root / 'page.pdf'), str(root / 'render'))
    assert Image.open(root / 'render.png').size == (200, 150)
    run('zip', '-q', '-j', str(root / 'archive.zip'), str(root / 'text.docx'))
    assert zipfile.ZipFile(root / 'archive.zip').namelist() == ['text.docx']
    assert run('awk', 'BEGIN {print 6*7}').strip() == '42'
    run('git', 'init', '--quiet', str(root / 'repository'))
    run('git', '-C', str(root / 'repository'), '-c', 'user.name=Fixture',
        '-c', 'user.email=fixture@example.invalid', 'commit', '--quiet', '--allow-empty', '-m', 'fixture')
    assert len(run('git', '-C', str(root / 'repository'), 'rev-parse', 'HEAD').strip()) == 40
print(json.dumps({'result': 'PASS', 'checks': ['image-font-webp', 'html-xml', 'docx', 'xlsx',
                                             'pdf-copy-render', 'zip', 'awk', 'git']}))
